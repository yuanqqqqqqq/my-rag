"""功能冒烟测试：把工具的各种能力和边界情况全跑一遍。

    python scripts/smoke_test.py

会临时造一批多格式语料（md / txt / py / js / docx / pdf），
建索引、检索、试错，最后清理掉临时索引。

与 tests/ 的区别：
    tests/       —— 纯函数，不联网，秒级，每次提交都跑（CI 里跑的就是它）
    smoke_test   —— 全链路，要联网和 API Key，几十秒，改了大改动后手动跑
"""

from __future__ import annotations

import shutil
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

TEST_INDEX = "smoke_test_tmp"

_results: list[tuple[str, bool, str]] = []


def check(name: str, ok: bool, detail: str = "") -> bool:
    _results.append((name, ok, detail))
    mark = "PASS" if ok else "FAIL"
    print(f"  [{mark}] {name}" + (f"  — {detail}" if detail else ""))
    return ok


# ══════════════════════════════════════════════════════
# 造一个最小的合法 PDF（不依赖任何 PDF 生成库）
# ══════════════════════════════════════════════════════


def make_pdf(path: Path, text: str) -> None:
    """手写一个单页 PDF，xref 偏移量按实际算 —— pypdf 要求它是对的。"""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n"
        + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"

    xref_pos = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF\n"
    ).encode()

    path.write_bytes(bytes(out))


def make_docx(path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading("产品需求文档", level=1)
    d.add_paragraph("本产品的目标是让用户在三步内完成部署。")
    d.add_heading("支持的环境", level=2)
    d.add_paragraph("支持 Linux 与 macOS，Windows 需要 WSL2。")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "目标值"
    table.cell(1, 0).text = "部署耗时"
    table.cell(1, 1).text = "小于 5 分钟"
    d.save(str(path))


def build_corpus(root: Path) -> None:
    """造一批覆盖所有支持格式的语料。"""
    (root / "指南.md").write_text(
        "# 部署指南\n\n## 环境准备\n\n需要 Python 3.10 以上版本，以及一个可用的 API Key。\n\n"
        "## 部署步骤\n\n第一步克隆代码仓库，第二步安装依赖，第三步配置环境变量。\n\n"
        "## 常见问题\n\n如果端口被占用，修改配置文件里的 port 字段。\n",
        encoding="utf-8",
    )
    (root / "notes.txt").write_text(
        "会议纪要：确定了三个关键里程碑，分别是原型、内测、公测。\n"
        "负责人分别是张三、李四、王五。\n",
        encoding="utf-8",
    )
    (root / "utils.py").write_text(
        "def calculate_retry_delay(attempt: int) -> float:\n"
        "    return min(2 ** attempt, 60.0)\n\n\n"
        "class RateLimiter:\n"
        "    \"\"\"令牌桶限流器。\"\"\"\n\n"
        "    def __init__(self, rate: int):\n"
        "        self.rate = rate\n",
        encoding="utf-8",
    )
    (root / "app.js").write_text(
        "export function parseConfig(raw) {\n  return JSON.parse(raw);\n}\n\n"
        "class EventBus {\n  constructor() { this.handlers = {}; }\n}\n",
        encoding="utf-8",
    )
    make_docx(root / "需求文档.docx")
    # 刻意让 PDF 的内容【和其他文档没有主题重叠】——
    # 冒烟测试断言的是精确名次，语料一旦有歧义，断言就变成掷骰子。
    # （第一版写的是 "deployment guide for Linux"，和需求文档里的
    #   「部署」「Linux」撞车，结果时对时错。）
    make_pdf(root / "release.pdf", "Release notes for version 3.2.1. This build fixes a memory leak.")

    # 不该被索引的噪声文件
    (root / "domain.txt").write_text("年假 100 n\n", encoding="utf-8")
    (root / "debug.log").write_text("x" * 100, encoding="utf-8")
    git = root / ".git"
    git.mkdir()
    (git / "config.md").write_text("# 不该被索引", encoding="utf-8")


# ══════════════════════════════════════════════════════


def main() -> int:
    from rag.client import get_client
    from rag.chunking import load_documents
    from rag.loaders import load_path
    from rag.retrieve import build_bm25, retrieve
    from rag.store import (
        build_collection,
        collection_info,
        drop_collection,
        get_collection,
        list_collections,
        load_chunks,
    )
    from rag.tokenize import tokenize

    print("=" * 74)
    print("my-rag 功能冒烟测试")
    print("=" * 74)

    tmp = Path(tempfile.mkdtemp(prefix="myrag_smoke_"))
    drop_collection(TEST_INDEX)  # 上次跑崩了可能残留，先清干净，保证"全量"断言成立
    try:
        # ── 1. 多格式加载 ──
        print("\n【1】文档加载")
        build_corpus(tmp)
        docs = load_path(tmp, exclude=["*.log"], verbose=False)
        kinds = {}
        for d in docs:
            kinds[d.kind] = kinds.get(d.kind, 0) + 1

        check("加载 markdown / text / code / docx / pdf 五类",
              all(kinds.get(k, 0) >= 1 for k in ("markdown", "text", "code", "docx", "pdf")),
              f"实际: {kinds}")
        check("自动跳过 .git 目录",
              not any(".git" in d.source for d in docs))
        check("--exclude 排除 .log",
              not any(d.source.endswith(".log") for d in docs))
        check("domain.txt 未被排除（没给这个规则）",
              any(d.source == "domain.txt" for d in docs))

        # ── 2. 切块 ──
        print("\n【2】切块")
        items = load_documents(str(tmp), None, exclude=["*.log"], verbose=False)
        by_kind: dict[str, int] = {}
        for it in items:
            by_kind[it["kind"]] = by_kind.get(it["kind"], 0) + 1
        check("切块覆盖所有类型", len(by_kind) >= 4, f"{by_kind}")
        check("每个块都有 source", all(it["source"] for it in items))
        check("每个块都有非空文本", all(it["text"].strip() for it in items))
        check("python 文件的块带 lang 标记",
              any(it.get("lang") == "python" for it in items))

        # ── 3. 建索引 ──
        print("\n【3】建索引")
        client = get_client()
        stats = build_collection(client, items, TEST_INDEX, verbose=False)
        col = get_collection(TEST_INDEX)
        check("索引块数与切块数一致", col.count() == len(items),
              f"{col.count()} = {len(items)}")
        check("首次建索引走全量模式", stats["mode"] == "full", stats["mode"])

        # ── 3b. 增量索引 ──
        stats2 = build_collection(client, items, TEST_INDEX, verbose=False)
        check("内容没变时，增量索引一块都不重算",
              stats2["added"] == 0 and stats2["reused"] == len(items),
              f"新增 {stats2['added']}，复用 {stats2['reused']}")

        # 改一个字，只应该重算那一个块
        mutated = [dict(it) for it in items]
        mutated[0]["text"] = mutated[0]["text"] + "\n\n新增一句话用于测试增量。"
        stats3 = build_collection(client, mutated, TEST_INDEX, verbose=False)
        check("改一个块只重算一个", stats3["added"] == 1 and stats3["reused"] == len(items) - 1,
              f"新增 {stats3['added']}，复用 {stats3['reused']}")
        check("旧版本的块被清掉", stats3["removed"] == 1, f"删除 {stats3['removed']}")

        # ── 4. 多索引共存 ──
        print("\n【4】多索引管理")
        names = {c["name"] for c in list_collections()}
        check("新索引出现在列表里", TEST_INDEX in names, f"共 {len(names)} 个索引")

        info = collection_info(TEST_INDEX)
        check("info 能报出类型分布", len(info["kinds"]) >= 4, f"{info['kinds']}")

        # ── 5. 检索 ──
        print("\n【5】检索（真实的 API 调用）")
        chunks = load_chunks(col)
        bm25 = build_bm25(chunks)

        # ⚠️ 每个查询必须【只有一个合理答案】。
        #
        #    第一版用了「部署的步骤是什么」，结果时对时错 ——
        #    因为 指南.md 和 需求文档.docx 里都出现了「部署」，
        #    谁排第一取决于 LLM 精排当次的心情。
        #
        #    这是个反复出现的规律：**候选同质时，精确名次没有意义，
        #    断言它只会让测试变得 flaky。**
        #    冒烟测试要的是确定性 —— 用只有一处能回答的问题。
        cases = [
            ("中文语义", "端口被占用了怎么办", "指南.md"),
            ("中文关键词", "会议纪要里提到了什么", "notes.txt"),
            ("英文 PDF", "which version fixed the memory leak", "release.pdf"),
            ("Word 正文", "支持哪些操作系统", "需求文档.docx"),
            ("代码", "calculate_retry_delay", "utils.py"),
            ("JS 代码", "parseConfig", "app.js"),
        ]

        for label, q, expect in cases:
            hits = retrieve(client, col, q, strategy="hybrid-rerank",
                            chunks=chunks, bm25=bm25)
            top = hits[0]["source"] if hits else "(空)"
            check(f"{label}: 命中 {expect}", top == expect, f"实际 Top1 = {top}")

        # ── 6. 中文分词 ──
        print("\n【6】中文分词")
        toks = tokenize("怎么申请年假")
        check("中文被切词", "申请" in toks, f"{toks}")
        check("单字未被丢弃（年/假都在）",
              ("年假" in toks) or ("年" in toks and "假" in toks), f"{toks}")

        # ── 7. 边界与错误处理 ──
        print("\n【7】边界情况")
        from rag.store import _validate_name

        try:
            _validate_name("ab")
            check("索引名过短应报错", False, "居然没报错")
        except ValueError as exc:
            check("索引名过短给出可操作的报错", "太短" in str(exc))

        try:
            get_collection("根本不存在_xyz")
            check("不存在的索引应报错", False, "居然没报错")
        except RuntimeError as exc:
            check("索引不存在时提示现有索引", "不存在" in str(exc))

        try:
            load_path(tmp / "no_such_dir", verbose=False)
            check("路径不存在应报错", False, "居然没报错")
        except FileNotFoundError:
            check("路径不存在报 FileNotFoundError", True)

        # ── 8. 配置文件 ──
        print("\n【8】配置文件（.ragignore / rag.toml / 参数持久化）")
        from rag.ingest import build_index
        from rag.project import load_project_config
        from rag.store import get_index_config

        cfg_dir = tmp / "cfgtest"
        cfg_dir.mkdir()
        (cfg_dir / "a.md").write_text("# A\n\n关于年假的内容。", encoding="utf-8")
        (cfg_dir / "debug.log").write_text("日志内容", encoding="utf-8")
        (cfg_dir / "dict.txt").write_text("年假 100 n\n", encoding="utf-8")
        (cfg_dir / ".ragignore").write_text("# 排除日志\n*.log\n", encoding="utf-8")
        (cfg_dir / "rag.toml").write_text(
            'name = "smoke_cfg"\ndict = "dict.txt"\n', encoding="utf-8"
        )

        pc = load_project_config(cfg_dir)
        check(".ragignore 被读取", "*.log" in pc.exclude, str(pc.exclude))
        check("rag.toml 被读取", pc.name == "smoke_cfg" and pc.dict_path == "dict.txt")

        stats_cfg = build_index([str(cfg_dir)], None, verbose=False)
        check("零参数建索引，索引名来自 rag.toml",
              stats_cfg["name"] == "smoke_cfg", stats_cfg["name"])
        check(".ragignore 生效（日志未被索引）",
              stats_cfg["chunks"] == 1, f"{stats_cfg['chunks']} 块")

        saved = get_index_config("smoke_cfg")
        check("分词词典被持久化到索引",
              bool(saved.get("dict")) and str(saved["dict"]).endswith("dict.txt"),
              str(saved.get("dict")))
        check("词典文件被自动排除在语料外",
              "dict.txt" in saved.get("exclude", []), str(saved.get("exclude")))
        check("工具自己的配置文件不会被索引",
              not any(s in ("rag.toml", ".ragignore")
                      for s in saved.get("source_paths", [])),
              str(saved.get("source_paths", ""))[:60])

        check("配置测试索引已清理", drop_collection("smoke_cfg"))

        # ── 9. 多轮对话 ──
        print("\n【9】多轮对话")
        from rag.conversation import rewrite_question
        from rag.rewrite import build_vocab, find_oov_terms

        vocab = build_vocab([c["text"] for c in chunks])

        first, warn = rewrite_question(client, "怎么申请年假", [])
        check("第一轮不改写（省一次 API 调用）",
              first == "怎么申请年假" and warn is None, f"-> {first}")

        history = [("怎么申请年假？", "员工需在 OA 系统提交年假申请单。")]
        rewritten, _ = rewrite_question(client, "那病假呢？", history)
        check("追问被补全指代",
              "病假" in rewritten and len(rewritten) > 4, f"-> {rewritten}")
        check("改写没引入语料外的词（无投毒）",
              not find_oov_terms(rewritten, vocab),
              str(find_oov_terms(rewritten, vocab)))

        # ── 10. HTTP 服务 ──
        print("\n【10】HTTP 服务")
        from fastapi.testclient import TestClient

        from rag.server import build_app

        http = TestClient(build_app())
        check("GET /health", http.get("/health").json() == {"status": "ok"})
        check("GET /indexes 返回列表",
              isinstance(http.get("/indexes").json(), list))
        home = http.get("/")
        check("GET / 返回网页", home.status_code == 200 and "<html" in home.text)

        r_ask = http.post("/ask", json={
            "question": "端口被占用了怎么办", "index": TEST_INDEX, "strategy": "vector",
        })
        check("POST /ask 返回答案",
              r_ask.status_code == 200 and "answer" in r_ask.json(),
              f"HTTP {r_ask.status_code}")

        r_bad = http.post("/ask", json={"question": "x", "index": "根本不存在的索引"})
        check("索引不存在时返回 400 而不是 500",
              r_bad.status_code == 400, f"HTTP {r_bad.status_code}")

        r_ret = http.post("/retrieve", json={
            "question": "端口", "index": TEST_INDEX, "strategy": "vector",
        })
        check("POST /retrieve 不生成答案",
              r_ret.status_code == 200 and "answer" not in r_ret.json())

        # ── 11. 清理 ──
        print("\n【11】清理")
        check("删除临时索引", drop_collection(TEST_INDEX))
        check("删除后不在列表里",
              TEST_INDEX not in {c["name"] for c in list_collections()})

    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    # ── 汇总 ──
    passed = sum(1 for _, ok, _ in _results if ok)
    total = len(_results)
    print("\n" + "=" * 74)
    print(f"结果：{passed}/{total} 通过")
    if passed < total:
        print("\n失败项：")
        for name, ok, detail in _results:
            if not ok:
                print(f"  - {name}  {detail}")
    print("=" * 74)
    return 1 if passed < total else 0


if __name__ == "__main__":
    sys.exit(main())
